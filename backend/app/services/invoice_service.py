"""
Order invoice generation as a Word (.docx) document.

Every order gets a printable invoice saved under uploads/invoices/ in .docx
format. The document includes the order items, totals, shipping address,
payment details and the full status history ("invoice messages"). The file is
regenerated when the order status or payment status changes so the stored
document always reflects the latest order state.
"""
from datetime import datetime, timezone
from pathlib import Path

from fastapi import HTTPException
from bson import ObjectId
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.core.database import get_db

INVOICE_DIR = Path(__file__).resolve().parents[3] / "uploads" / "invoices"
BRAND = "BD Shop"
COMPANY_NAME = "Bangladesh E-Commerce + ERP"
COMPANY_ADDRESS = "Head Office, Dhaka 1207, Bangladesh"
SUPPORT = "support@bdshop.local  .  +880 1700-000000"

PRIMARY = RGBColor(0x0F, 0x17, 0x2A)
ACCENT = RGBColor(0x71, 0x63, 0xFF)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def _fmt_money(value) -> str:
    try:
        return f"BDT {float(value or 0):,.2f}"
    except (TypeError, ValueError):
        return "BDT 0.00"


def _fmt_dt(value) -> str:
    if not value:
        return ""
    return value.strftime("%d %b %Y, %I:%M %p") if hasattr(value, "strftime") else str(value)


def _set_cell(cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None, size: int = 10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return cell


def _status_styling(doc: Document) -> None:
    """Light global styling so the invoice looks clean in Word."""
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)


async def generate_order_invoice(order_id: str) -> Path:
    """Build (or rebuild) the .docx invoice for an order and store it on disk."""
    db = get_db()
    try:
        order = await db.orders.find_one({"_id": ObjectId(order_id)})
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid order id")
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    order_number = order.get("order_number") or str(order["_id"])
    customer = None
    if order.get("customer_id"):
        customer = await db.customers.find_one({"_id": order["customer_id"]})

    doc = Document()
    _status_styling(doc)
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # ----- Header -----
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(BRAND)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = PRIMARY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(COMPANY_NAME)
    run.font.size = Pt(11)
    run.font.color.rgb = ACCENT

    addr = doc.add_paragraph()
    addr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = addr.add_run(f"{COMPANY_ADDRESS}\n{SUPPORT}")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("INVOICE")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = ACCENT

    # ----- Meta rows -----
    meta = doc.add_table(rows=0, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    status = order.get("status") or "unknown"
    payment_status = order.get("payment_status") or "pending"
    meta_rows = [
        ("Invoice No.", f"INV-{order_number}"),
        ("Order No.", order_number),
        ("Order Date", _fmt_dt(order.get("created_at"))),
        ("Order Status", status),
        ("Payment Method", order.get("payment_method") or "—"),
        ("Payment Status", payment_status),
        ("Payment Transaction", order.get("payment_transaction_id") or "—"),
    ]
    for label, value in meta_rows:
        row = meta.add_row().cells
        _set_cell(row[0], label, bold=True, color=PRIMARY)
        _set_cell(row[1], value)
    doc.add_paragraph()

    # ----- Bill To -----
    bill = doc.add_paragraph()
    run = bill.add_run("Bill To")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = PRIMARY

    bt = doc.add_table(rows=0, cols=2)
    bt.style = "Table Grid"
    ship = order.get("shipping_address") or {}
    bt_rows = [
        ("Customer", (customer or {}).get("full_name") or ship.get("full_name") or "—"),
        ("Email", (customer or {}).get("email") or "—"),
        ("Phone", (customer or {}).get("phone") or ship.get("phone") or "—"),
        ("Division", ship.get("division") or "—"),
        ("District / Upazila", f"{ship.get('district') or '—'} / {ship.get('upazila') or '—'}"),
        ("Address", ship.get("address_line") or "—"),
        ("Postal Code", ship.get("postal_code") or "—"),
    ]
    for label, value in bt_rows:
        row = bt.add_row().cells
        _set_cell(row[0], label, bold=True, color=PRIMARY)
        _set_cell(row[1], value)
    doc.add_paragraph()

    # ----- Items table -----
    items = doc.add_table(rows=1, cols=6)
    items.style = "Table Grid"
    items.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(0.3), Inches(2.6), Inches(1.1), Inches(0.6), Inches(1.0), Inches(1.1)]
    header = items.rows[0].cells
    for i, text in enumerate(["#", "Item", "SKU", "Qty", "Unit Price", "Total"]):
        _set_cell(header[i], text, bold=True, color=PRIMARY)

    for idx, it in enumerate(order.get("items") or [], start=1):
        row = items.add_row().cells
        row[0].width = widths[0]
        row[1].width = widths[1]
        row[2].width = widths[2]
        row[3].width = widths[3]
        row[4].width = widths[4]
        row[5].width = widths[5]
        _set_cell(row[0], str(idx))
        _set_cell(row[1], it.get("name") or "—")
        _set_cell(row[2], it.get("sku") or "—")
        _set_cell(row[3], str(it.get("quantity") or 0))
        _set_cell(row[4], _fmt_money(it.get("unit_price")))
        _set_cell(row[5], _fmt_money(it.get("line_total")), align=WD_ALIGN_PARAGRAPH.RIGHT)

    # ----- Totals -----
    doc.add_paragraph()
    totals = doc.add_table(rows=0, cols=2)
    totals.alignment = WD_TABLE_ALIGNMENT.RIGHT
    totals.style = "Table Grid"
    total_rows = [
        ("Subtotal", _fmt_money(order.get("subtotal"))),
        ("Discount", f"- {_fmt_money(order.get('discount'))}"),
        ("Delivery Charge", _fmt_money(order.get("delivery_charge"))),
        ("Tax", _fmt_money(order.get("tax"))),
        ("GRAND TOTAL", _fmt_money(order.get("total"))),
    ]
    for i, (label, value) in enumerate(total_rows):
        row = totals.add_row().cells
        bold = i == len(total_rows) - 1
        size = 11 if bold else 10
        _set_cell(row[0], label, bold=bold, color=PRIMARY, size=size)
        _set_cell(row[1], value, bold=bold, align=WD_ALIGN_PARAGRAPH.RIGHT, color=ACCENT, size=size)
    doc.add_paragraph()

    # ----- Status history ("invoice messages") -----
    hist = doc.add_paragraph()
    run = hist.add_run("Order Status History")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = PRIMARY

    history = order.get("history") or []
    if not history:
        p = doc.add_paragraph()
        run = p.add_run(f"• Placed as {status}")
        run.font.size = Pt(10)
    for entry in history:
        p = doc.add_paragraph()
        run = p.add_run(
            f"• {entry.get('status') or status} — {_fmt_dt(entry.get('at'))}"
            f"{' (by ' + str(entry.get('by')) + ')' if entry.get('by') else ''}"
        )
        run.font.size = Pt(10)

    # Notes
    if order.get("notes"):
        doc.add_paragraph()
        p = doc.add_paragraph()
        run = p.add_run(f"Order Notes: {order['notes']}")
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED

    # ----- Footer -----
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run(f"Generated {_fmt_dt(datetime.now(timezone.utc))}  ·  {BRAND}  ·  {COMPANY_NAME}")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    INVOICE_DIR.mkdir(parents=True, exist_ok=True)
    path = INVOICE_DIR / f"INV-{order_number}.docx"
    doc.save(str(path))

    await db.orders.update_one(
        {"_id": order["_id"]},
        {
            "$set": {
                "invoice_path": str(path),
                "invoice_generated_at": datetime.now(timezone.utc),
            }
        },
    )
    return path


async def ensure_order_invoice(order_id: str) -> Path:
    """Return the existing invoice .docx for an order, generating it if needed."""
    db = get_db()
    try:
        order = await db.orders.find_one({"_id": ObjectId(order_id)})
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid order id")
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    stored = order.get("invoice_path")
    if stored and Path(stored).exists():
        return Path(stored)
    return await generate_order_invoice(order_id)


async def regenerate_order_invoice(order_id: str) -> None:
    """Safety wrapper for regenerating an invoice after a status change."""
    try:
        await generate_order_invoice(order_id)
    except Exception:
        pass